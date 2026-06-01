"""报名导出业务逻辑层."""

import csv
import io
from typing import Optional

from django.http import HttpResponse

from api.models import StudentApplication


def export_applications(
    status: Optional[str] = None,
    format: str = "csv",
) -> HttpResponse:
    """导出报名数据.

    Args:
        status: 状态筛选
        format: 导出格式 (csv/xlsx)

    Returns:
        HTTP 响应对象
    """
    queryset = StudentApplication.objects.all()

    if status and status != "all":
        queryset = queryset.filter(status=status)

    applications = list(queryset.order_by("-created_at"))

    if format == "xlsx":
        return _export_xlsx(applications)
    else:
        return _export_csv(applications)


def _export_csv(applications: list[StudentApplication]) -> HttpResponse:
    """导出为 CSV 格式.

    Args:
        applications: 报名列表

    Returns:
        HTTP 响应对象
    """
    response = HttpResponse(
        content_type="text/csv; charset=utf-8-sig",
    )
    response["Content-Disposition"] = 'attachment; filename="applications.csv"'

    writer = csv.writer(response)
    writer.writerow(
        [
            "学号",
            "姓名",
            "性别",
            "年级",
            "专业",
            "手机号",
            "邮箱",
            "GPA",
            "技能",
            "项目经验",
            "报名动机",
            "状态",
            "提交时间",
            "审核时间",
            "审核意见",
        ]
    )

    gender_map = {"male": "男", "female": "女"}
    status_map = dict(StudentApplication.STATUS_CHOICES)

    for app in applications:
        writer.writerow(
            [
                app.student_id,
                app.name,
                gender_map.get(app.gender, app.gender),
                app.grade,
                app.major,
                app.phone,
                app.email,
                str(app.gpa) if app.gpa else "",
                ", ".join(app.skills) if app.skills else "",
                app.experience,
                app.motivation,
                status_map.get(app.status, app.status),
                app.submitted_at.strftime("%Y-%m-%d %H:%M") if app.submitted_at else "",
                app.reviewed_at.strftime("%Y-%m-%d %H:%M") if app.reviewed_at else "",
                app.review_comment,
            ]
        )

    return response


def _export_xlsx(applications: list[StudentApplication]) -> HttpResponse:
    """导出为 Excel 格式.

    Args:
        applications: 报名列表

    Returns:
        HTTP 响应对象
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as err:
        raise RuntimeError(
            "请安装 openpyxl: pip install openpyxl 或 uv sync --extra excel"
        ) from err

    wb = Workbook()
    ws = wb.active
    ws.title = "报名数据"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="4472C4",
        end_color="4472C4",
        fill_type="solid",
    )
    header_alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        "学号",
        "姓名",
        "性别",
        "年级",
        "专业",
        "手机号",
        "邮箱",
        "GPA",
        "技能",
        "项目经验",
        "报名动机",
        "状态",
        "提交时间",
        "审核时间",
        "审核意见",
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    gender_map = {"male": "男", "female": "女"}
    status_map = dict(StudentApplication.STATUS_CHOICES)

    for row_idx, app in enumerate(applications, 2):
        ws.cell(row=row_idx, column=1, value=app.student_id)
        ws.cell(row=row_idx, column=2, value=app.name)
        ws.cell(row=row_idx, column=3, value=gender_map.get(app.gender, app.gender))
        ws.cell(row=row_idx, column=4, value=app.grade)
        ws.cell(row=row_idx, column=5, value=app.major)
        ws.cell(row=row_idx, column=6, value=app.phone)
        ws.cell(row=row_idx, column=7, value=app.email)
        ws.cell(row=row_idx, column=8, value=str(app.gpa) if app.gpa else "")
        ws.cell(row=row_idx, column=9, value=", ".join(app.skills) if app.skills else "")
        ws.cell(row=row_idx, column=10, value=app.experience)
        ws.cell(row=row_idx, column=11, value=app.motivation)
        ws.cell(row=row_idx, column=12, value=status_map.get(app.status, app.status))
        ws.cell(
            row=row_idx,
            column=13,
            value=app.submitted_at.strftime("%Y-%m-%d %H:%M") if app.submitted_at else "",
        )
        ws.cell(
            row=row_idx,
            column=14,
            value=app.reviewed_at.strftime("%Y-%m-%d %H:%M") if app.reviewed_at else "",
        )
        ws.cell(row=row_idx, column=15, value=app.review_comment)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="applications.xlsx"'

    return response
