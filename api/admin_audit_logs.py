"""管理员后台 — 审计日志模块路由."""

import io

from django.http import HttpRequest, StreamingHttpResponse
from ninja import Router
from ninja.pagination import LimitOffsetPagination, paginate

from api.admin_schemas import AuditLogOut, ForbiddenResponse
from api.dependencies import permission_required
from api.models import AuditLog

router = Router(tags=["Admin-AuditLogs"])

SENSITIVE_FIELDS = {"password", "token", "secret", "old_password", "new_password", "refresh_token", "access_token"}


def _mask_sensitive_data(data: dict | None) -> dict | None:
    """对敏感字段进行掩码处理.

    递归遍历字典，将包含密码、token 等敏感信息的字段值替换为 ***。

    Args:
        data: 原始数据字典

    Returns:
        掩码后的数据字典
    """
    if data is None:
        return None

    if not isinstance(data, dict):
        return data

    masked = {}
    for key, value in data.items():
        if any(s in key.lower() for s in SENSITIVE_FIELDS):
            masked[key] = "***"
        elif isinstance(value, dict):
            masked[key] = _mask_sensitive_data(value)
        elif isinstance(value, list):
            masked[key] = [
                _mask_sensitive_data(item) if isinstance(item, dict) else item
                for item in value
            ]
        else:
            masked[key] = value
    return masked


def _log_to_out(log: AuditLog) -> dict:
    return {
        "id": log.id,
        "user_id": log.user_id,
        "user_name": log.user.username if log.user else "未知用户",
        "action": log.action,
        "action_display": log.get_action_display(),
        "module": log.module,
        "target_id": log.target_id,
        "target_type": log.target_type,
        "ip_address": str(log.ip_address) if log.ip_address else None,
        "details": _mask_sensitive_data(log.details),
        "created_at": log.created_at,
    }


@router.get(
    "",
    response={200: list[AuditLogOut], 403: ForbiddenResponse},
    summary="获取审计日志列表",
    description="获取审计日志分页列表，支持按操作人、模块、时间范围、Action 类型筛选。",
)
@paginate(LimitOffsetPagination)
@permission_required("system.manage")
def list_audit_logs(
    request: HttpRequest,
    user_id: int | None = None,
    module: str | None = None,
    action: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
) -> list[dict]:
    """获取审计日志列表."""
    queryset = AuditLog.objects.all().select_related("user")

    if user_id is not None:
        queryset = queryset.filter(user_id=user_id)

    if module:
        queryset = queryset.filter(module__icontains=module)

    if action:
        queryset = queryset.filter(action=action.upper())

    if date_from:
        queryset = queryset.filter(created_at__date__gte=date_from)

    if date_to:
        queryset = queryset.filter(created_at__date__lte=date_to)

    return [_log_to_out(log) for log in queryset]


@router.get(
    "/export",
    summary="导出审计日志 Excel",
    description="导出审计日志报表为 Excel 文件（xlsx），对敏感字段进行掩码处理。",
)
@permission_required("system.manage")
def export_audit_logs(request: HttpRequest) -> StreamingHttpResponse:
    """导出审计日志 Excel.

    使用 openpyxl 生成 xlsx 文件。
    导出时对 details 中的密码/token 等敏感字段进行掩码处理。
    """
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "审计日志"

    headers = ["ID", "操作人", "操作类型", "操作模块", "目标ID", "目标类型", "IP地址", "操作详情", "操作时间"]
    header_font = Font(bold=True)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font

    row_num = 2
    for log in AuditLog.objects.all().select_related("user").iterator():
        ws.cell(row=row_num, column=1, value=log.id)
        ws.cell(row=row_num, column=2, value=log.user.username if log.user else "未知用户")
        ws.cell(row=row_num, column=3, value=log.get_action_display())
        ws.cell(row=row_num, column=4, value=log.module)
        ws.cell(row=row_num, column=5, value=log.target_id or "")
        ws.cell(row=row_num, column=6, value=log.target_type or "")
        ws.cell(row=row_num, column=7, value=str(log.ip_address) if log.ip_address else "")
        ws.cell(row=row_num, column=8, value=str(_mask_sensitive_data(log.details)))
        ws.cell(row=row_num, column=9, value=log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else "")
        row_num += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = StreamingHttpResponse(
        output,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="audit_logs.xlsx"'
    return response
